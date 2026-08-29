from __future__ import annotations

import shutil
import zipfile
from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\kiadt\Desktop\shop\longtua-apartment")
REFERENCE = Path(r"C:\Users\kiadt\.codex\skills\artifact-template-system-design\assets\reference.docx")
WORK = ROOT / ".artifacts" / "system-design-work"
OUTPUT = ROOT / "docs" / "longtua-shared-ui-system-design.docx"
DIAGRAM = WORK / "architecture.png"


def set_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph(paragraph, text)
    for extra in cell.paragraphs[1:]:
        set_paragraph(extra, "")


def fill_table(table, rows: list[list[str]]) -> None:
    while len(table.rows) < len(rows):
        source = table.rows[-1]
        new_tr = deepcopy(source._tr)
        table._tbl.append(new_tr)
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            set_cell(cell, value)


def font(size: int, bold: bool = False):
    path = Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf")
    return ImageFont.truetype(str(path), size=size)


def draw_diagram() -> None:
    image = Image.new("RGBA", (1568, 800), "white")
    draw = ImageDraw.Draw(image)
    navy, blue, pale, border, text = "#102A43", "#2563EB", "#EAF2FF", "#B8C9DB", "#334E68"

    def box(x1, y1, x2, y2, title, lines, accent=False):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=pale if accent else "#F7FAFC", outline=blue if accent else border, width=4)
        draw.text((x1 + 24, y1 + 18), title, fill=navy, font=font(30, True))
        for index, line in enumerate(lines):
            draw.text((x1 + 24, y1 + 64 + index * 34), line, fill=text, font=font(23))

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill=blue, width=6)
        draw.polygon([(x2, y2), (x2 - 18, y2 - 11), (x2 - 18, y2 + 11)], fill=blue)

    draw.text((48, 28), "Longtua shared UI with isolated security boundaries", fill=navy, font=font(38, True))
    box(55, 115, 440, 335, "Owner Portal", ["Static feature routes", "Organization-scoped reads", "Customer server actions"], True)
    box(55, 465, 440, 685, "Admin Console", ["Explicit admin routes", "Cross-organization reads", "Admin-only server actions"], True)
    box(590, 245, 990, 555, "Shared presentation", ["DataTable + EmptyState", "StatusBadge", "Select + Date/Time", "Thai formatters", "Feature composition"], False)
    box(1140, 115, 1510, 335, "Portal loaders", ["Authenticated session", "Membership + permissions", "Serializable view models"], False)
    box(1140, 465, 1510, 685, "Admin loaders", ["System-admin check", "Section-specific queries", "Serializable view models"], False)
    arrow(440, 225, 590, 330)
    arrow(440, 575, 590, 470)
    arrow(990, 330, 1140, 225)
    arrow(990, 470, 1140, 575)
    draw.text((580, 730), "Supabase: RLS for tenant data | service role remains server-only", fill=navy, font=font(25, True))
    image.save(DIAGRAM)


def replace_media(docx_path: Path) -> None:
    temp = docx_path.with_suffix(".media.docx")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = DIAGRAM.read_bytes() if item.filename == "word/media/image1.png" else source.read(item.filename)
            target.writestr(item, payload)
    temp.replace(docx_path)


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    working = WORK / "working.docx"
    shutil.copy2(REFERENCE, working)
    doc = Document(working)

    paragraphs = doc.paragraphs
    replacements = {
        8: "Longtua Apartment Management",
        9: "Shared UI and Security Boundary Architecture",
        22: "This proposal removes duplicated presentation code between the owner portal and Super Admin console while preserving separate authorization and data-access boundaries. Both surfaces use one set of tables, status badges, empty states, date/select controls, and Thai formatting utilities. Server Components continue to read data directly, while mutations remain role-specific Server Actions.",
        23: "The design supports multi-organization apartment operations and platform administration. It does not merge the two layouts, expose service-role credentials, or allow Admin behavior to bypass portal permissions. Admin queries are selected by page so a route does not load unrelated operational tables.",
        28: "Previously, the portal defined DataTable, StatusBadge, EmptyState, and formatting behavior in portal-specific files, while Admin implemented parallel AdminTable, status labels, and date/money helpers. The Admin section loader also queried most platform tables for every route. This increased maintenance cost, visual drift, and database work.",
        29: "The system now separates route and security boundaries from presentation primitives. Route pages and loaders determine audience, authorization, organization scope, and available commands. Shared UI components render serializable view models without knowing whether data came from a tenant-scoped or platform-scoped query.",
        33: "Figure 1. Shared presentation components with isolated Portal and Admin security/data boundaries.",
        40: "1. A request enters through an explicit Portal route such as /guestrooms or Admin route such as /admin/rooms.",
        41: "2. The Server Component validates the authenticated session, then checks organization membership or system-admin status. Authorization fails closed.",
        42: "3. The page loads only the configuration and operational records required by its feature. Portal reads remain organization-scoped; Admin reads use the server-only admin client.",
        43: "4. The loader creates plain serializable values and lookup maps for server-side composition. Shared UI receives strings, numbers, arrays, and React nodes rather than database clients or credentials.",
        44: "5. A mutation is accepted only by the audience-specific Server Action, which repeats authorization and validates identifiers and form values before writing.",
        45: "6. Database constraints and RLS protect tenant data. Administrative writes use explicit audit logging; destructive operations fail when referenced billing or lease history exists.",
        46: "7. Successful mutations revalidate affected routes. Failures return actionable messages while server logs retain request identifiers and sanitized diagnostics.",
        54: "All Client Component props crossing the server boundary must be JSON-serializable, except explicitly marked Server Actions.",
        55: "Mutations identify organization and entity records with UUIDs and repeat authorization inside the Server Action.",
        56: "Billing records retain rate snapshots, period identifiers, meter readings, and audit identifiers needed to explain historical totals.",
        57: "Shared UI is a presentation contract only; Supabase tables and server-side authorization remain the source of truth.",
        58: "The implementation contract is represented by components/ui, lib/format.ts, explicit App Router pages, and architecture tests in the repository.",
        59: "Contract changes require lint, automated tests, TypeScript production build, and review of both Portal and Admin consumers.",
        63: "Server Actions validate each request independently, so browser retries cannot elevate privileges. Existing database uniqueness constraints prevent duplicate room, invoice, and document identifiers. Revalidation occurs only after a successful durable write. Audit records preserve administrative changes; billing snapshots preserve the values used at issuance time.",
        67: "Portal access requires an active organization membership and granular menu/action permission. Admin access requires an active system_admins record.",
        68: "Client props contain only display data required by the page. Passwords, service-role keys, database clients, and internal errors never cross the Server Component boundary.",
        69: "Supabase service-role credentials are imported only by server-only modules. Customer operations use the authenticated Supabase client and RLS.",
        70: "Unknown permissions, missing memberships, invalid UUIDs, disabled subscriptions, and unavailable schemas fail closed. No fallback organization or demo data is used.",
        71: "Deletion is blocked by foreign-key references to leases, invoices, payments, and meter history. Administrative changes are written to platform audit logs.",
        81: "Should lists above the current server limits adopt cursor pagination or virtualization first?",
        82: "Which domain components should move next into components/features after usage patterns stabilize?",
        83: "Should production visual-regression coverage run for both Portal and Admin themes in CI?",
        84: "Which team owns SLO thresholds and alert routing for admin cross-organization queries?",
        87: "Adopt the shared-presentation/separate-boundary architecture. The first milestone is complete: common table, status, empty-state, control, and formatting code is shared; routes remain explicit; Admin queries are section-selective. Next, add browser visual regression and introduce pagination where measured data volume requires it.",
    }
    for index, value in replacements.items():
        set_paragraph(paragraphs[index], value)
    # This template slot contains a hyperlink run that is not exposed through
    # paragraph.runs. Clear that single slot so no retained placeholder remains.
    paragraphs[59].clear()
    paragraphs[59].add_run(replacements[59])

    fill_table(doc.tables[0], [["STATUS\nApproved", "", "OWNER\nLongtua Engineering", "", f"LAST UPDATED\n{date.today():%B %d, %Y}"]])
    fill_table(doc.tables[1], [
        ["Authors", "Longtua Engineering"],
        ["Reviewers", "Product owner and platform security reviewer"],
        ["Related docs", "README.md; tests/portal-architecture.test.mjs"],
        ["Scope", "Shared UI architecture for owner Portal and Super Admin while preserving authorization and data isolation."],
    ])
    fill_table(doc.tables[2], [
        ["Goals", "Non-goals"],
        ["One implementation for tables, badges, empty states, controls, and formatters.", "A single shared full-page component for Portal and Admin."],
        ["Keep tenant and platform authorization boundaries independent and fail-closed.", "Changing the Supabase schema or weakening RLS."],
        ["Use explicit routes and section-specific Admin queries.", "Rebuilding all feature UX or visual branding."],
        ["Keep changes verifiable through lint, tests, and production build.", "Adding pagination without measured scale requirements."],
    ])
    fill_table(doc.tables[3], [
        ["Component", "Responsibility", "Primary storage", "Failure behavior"],
        ["Portal routes/loaders", "Authenticate members, scope by organization, prepare view models.", "Supabase via authenticated client", "Redirect or deny; never fall back to another organization."],
        ["Admin routes/loaders", "Authorize system admins and query only page dependencies.", "Supabase via server-only admin client", "Redirect non-admins; show unavailable-schema state."],
        ["Shared UI", "Render DataTable, EmptyState, StatusBadge, select/date controls.", "No durable storage", "Render deterministic empty or validation states."],
        ["Audience Server Actions", "Validate, reauthorize, mutate, audit, and revalidate.", "Supabase operational and audit tables", "Return actionable error; no partial privileged fallback."],
        ["Architecture tests", "Prevent route, RSC, authorization, and UI-sharing regressions.", "Repository test suite", "Block merge/build when contracts drift."],
    ])
    fill_table(doc.tables[4], [
        ["Field", "Type", "Required", "Description"],
        ["section", "AdminSection", "Yes", "Compile-time route/view key supplied by an explicit Admin page."],
        ["organizationId", "UUID string", "Portal writes", "Selected organization; revalidated against active membership."],
        ["headers", "string[]", "Yes", "Column labels for the shared DataTable."],
        ["rows", "ReactNode[][]", "Yes", "Server- or client-composed serializable presentation cells."],
        ["status", "string", "Status views", "Canonical domain status mapped by shared statusLabel."],
        ["compact", "boolean", "No", "Density variant for Admin tables/badges/empty states."],
        ["searchParams", "Promise<object>", "Route dependent", "Next.js request-time filters and operation notices."],
    ])
    fill_table(doc.tables[5], [
        ["Scenario", "Expected behavior", "Reasoning"],
        ["Repeated form submission", "Constraints reject duplicates or action returns stable validation feedback.", "Durable identifiers and uniqueness remain server-controlled."],
        ["Missing membership or permission", "Request redirects or returns a denial without querying another tenant.", "Authorization is repeated at every server boundary."],
        ["Admin dependency unavailable", "Only the affected page shows an unavailable-data state.", "Section-selective queries reduce unrelated failure impact."],
        ["UI component changes", "Portal and Admin adopt the same implementation after tests/build pass.", "Presentation code has one source of truth."],
    ])
    fill_table(doc.tables[6], [
        ["Signal", "SLO or alert", "Owner", "Launch gate"],
        ["Authorization denials", "Monitor unexpected increases by route and audience.", "Platform security", "Required"],
        ["Admin page latency", "Track p95 per explicit section; investigate query-limit saturation.", "Platform engineering", "Required"],
        ["Server Action failures", "Track error rate and sanitized request IDs by action.", "Feature owner", "Required"],
        ["UI regression", "Lint, tests, production build, and browser visual checks.", "Frontend engineering", "Required"],
        ["Tenant isolation", "Automated RLS and organization-scope tests remain green.", "Platform security", "Required"],
        ["Rollout constraint: deploy shared primitives and query selection together; retain rollback through version control and validate Portal/Admin smoke paths before promotion.", "", "", ""],
    ])
    fill_table(doc.tables[7], [
        ["Alternative", "Why it was considered", "Why it was not selected"],
        ["One full page for both audiences", "Maximum superficial reuse.", "Mixes authorization, data scope, and privileged controls."],
        ["Separate duplicate UI trees", "Simple local ownership.", "Creates drift and doubles fixes for tables, status, and empty states."],
        ["Dynamic /admin/[section]", "Minimal route files.", "Known sections lose explicit ownership and route-level clarity."],
        ["Client-side API fetching", "Could centralize reads behind endpoints.", "Adds waterfalls and an unnecessary internal API layer."],
    ])
    fill_table(doc.tables[8], [
        ["Milestone", "Deliverable", "Exit criteria"],
        ["M1", "Shared primitives and centralized Thai formatters", "Portal/Admin consumers use the same implementation; tests pass."],
        ["M2", "Explicit routes and section-selective Admin queries", "No dynamic section route; production build lists each page."],
        ["M3", "Browser visual regression and performance baselines", "Critical Portal/Admin paths verified at desktop and mobile widths."],
        ["M4", "Pagination/feature extraction based on measurements", "Volume thresholds, SLOs, and rollback plan approved."],
    ])

    draw_diagram()
    doc.save(OUTPUT)
    replace_media(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
